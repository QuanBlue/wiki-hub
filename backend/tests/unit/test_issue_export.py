"""The HTML an issue's Word export is built from."""

from __future__ import annotations

import io
import uuid
from datetime import UTC, datetime
from pathlib import Path
from typing import Any
from unittest.mock import AsyncMock, Mock, patch

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml.parser import OxmlElement
from docx.shared import RGBColor

from app.core.config import settings
from app.modules.issues import export as export_module
from app.modules.issues import export_style
from app.modules.issues.export import (
    build_issue_html,
    build_issues_html,
    issue_to_docx,
    issues_to_docx,
    load_issue_images,
)
from app.modules.issues.export_style import (
    PRIMARY,
    PRIMARY_STRONG,
    STATUS_COLORS,
    SUNKEN,
    label_style,
    status_style,
    style_issue_report,
)
from app.schemas.issue import IssueNoteRead, IssuePerson, IssueRead


def _issue(**overrides: object) -> IssueRead:
    reporter = IssuePerson(id=uuid.uuid4(), username="alice", full_name="Alice Nguyen")
    values: dict[str, object] = {
        "id": uuid.uuid4(),
        "title": "Save fails",
        "description": "Steps:\n1. Open **page**\n<script>alert(1)</script>\n![x](http://evil.test/a.png)",
        "status": "in_progress",
        "page_url": "/spaces/ENG/pages/home",
        "labels": ["bug"],
        "reporter": reporter,
        "assignee": None,
        "notes": [
            IssueNoteRead(
                id=uuid.uuid4(),
                author=reporter,
                body="Public closing note",
                public=True,
                created_at=datetime(2026, 1, 2, tzinfo=UTC),
            ),
            IssueNoteRead(
                id=uuid.uuid4(),
                author=reporter,
                body="Internal triage note",
                public=False,
                created_at=datetime(2026, 1, 2, tzinfo=UTC),
            ),
        ],
        "created_at": datetime(2026, 1, 1, 9, 30, tzinfo=UTC),
        "updated_at": datetime(2026, 1, 2, tzinfo=UTC),
    }
    values.update(overrides)
    return IssueRead(**values)  # type: ignore[arg-type]


def test_issue_html_renders_markdown_safely_and_skips_internal_notes() -> None:
    html = build_issue_html(_issue(), [("shot.png", "image/png", b"\x89PNG")], lang="en")

    assert "<strong>page</strong>" in html
    assert "<script>" not in html
    assert "&lt;script&gt;" in html
    # A Markdown image would make pandoc download the URL - never rendered.
    assert '<img src="http' not in html
    assert "data:image/png;base64," in html
    assert "Public closing note" in html
    assert "Internal triage note" not in html
    assert "Alice Nguyen (@alice)" in html
    assert "In progress" in html and "Bug" in html
    assert "Unassigned" in html
    assert "Jan 1, 2026 09:30 UTC" in html
    assert 'custom-style="WH Status in_progress"' in html
    assert 'custom-style="WH Label bug"' in html


def test_issue_html_uses_vietnamese_labels() -> None:
    html = build_issue_html(_issue(description=""), [], lang="vi")

    assert "Mô tả" in html
    assert "Không có mô tả." in html
    assert "Đang xử lý" in html
    assert "Chưa giao" in html


def test_many_issues_get_a_summary_and_one_chapter_each() -> None:
    first = _issue(title="First problem")
    second = _issue(title="Second <problem>", status="done")
    html = build_issues_html([(first, []), (second, [])], lang="en")

    assert html.index("<h2>Summary</h2>") < html.index("<h1>First problem</h1>")
    assert html.index("<h1>First problem</h1>") < html.index("<h1>Second &lt;problem&gt;</h1>")
    # Each issue opens with its number, sections one level down.
    assert html.count('custom-style="WH Issue Eyebrow"') == 2
    assert "Issue 2 of 2" in html
    assert html.count("<h2>Details</h2>") == 2
    assert 'custom-style="WH Status done"' in html


def test_issue_html_covers_the_optional_rows_and_note_spacing(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "public_url", "https://wiki.example.com/")
    reporter = IssuePerson(id=uuid.uuid4(), username="bob", full_name="")
    notes = [
        IssueNoteRead(
            id=uuid.uuid4(),
            author=reporter,
            body=body,
            public=True,
            created_at=datetime(2026, 1, 2, tzinfo=UTC),
        )
        for body in ("First line\nsecond line", "Another note")
    ]
    html = build_issue_html(
        _issue(
            labels=["bug", "custom-tag"],
            resolved_at=datetime(2026, 1, 3, tzinfo=UTC),
            notes=notes,
        ),
        [],
        lang="en",
    )

    # A relative page path is made absolute against the public URL and linked.
    assert '<a href="https://wiki.example.com/spaces/ENG/pages/home">' in html
    assert "Resolved" in html and "Jan 3, 2026" in html
    # A slug with no colour of its own is printed plainly, not as a pill.
    assert "custom-tag" in html
    assert 'custom-style="WH Label custom-tag"' not in html
    # Notes after the first are spaced apart; a name-less author is @handle.
    assert html.count('custom-style="WH Gap"') == 1
    assert "@bob · " in html
    assert "First line<br />second line" in html


def test_issue_html_keeps_a_page_url_it_cannot_link(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(settings, "public_url", "")
    html = build_issue_html(
        _issue(description="", page_url="/spaces/ENG/pages/home"), [], lang="en"
    )

    assert "<a href=" not in html
    assert "/spaces/ENG/pages/home" in html


def _pandoc_like_docx() -> bytes:
    """What pandoc hands back for a report: a stock Title paragraph, the
    ``custom-style`` paragraphs, a details table and a summary table."""
    document = Document()
    document.add_paragraph("Report", style="Title")
    for name in ("WH Meta Label", "WH Meta Value", "WH Table Head", "WH Table Cell"):
        document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    document.add_heading("Save fails", level=1)

    details = document.add_table(rows=1, cols=2)
    details.cell(0, 0).paragraphs[0].style = "WH Meta Label"
    details.cell(0, 0).paragraphs[0].add_run("Assignee")
    details.cell(0, 1).paragraphs[0].style = "WH Meta Value"

    summary = document.add_table(rows=2, cols=5)
    for cell in summary.rows[0].cells:
        cell.paragraphs[0].style = "WH Table Head"
    for cell in summary.rows[1].cells:
        cell.paragraphs[0].style = "WH Table Cell"

    document.sections[0].footer.paragraphs[0].add_run("stale footer")

    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def _load(data: bytes) -> Any:
    return Document(io.BytesIO(data))


@pytest.mark.parametrize("multi", [False, True])
def test_style_issue_report_paints_the_wiki_workspace_look(multi: bool) -> None:
    document = _load(_pandoc_like_docx())
    style_issue_report(document, multi=multi, footer_text="WikiHub issue report")

    # pandoc's own title paragraph is gone - the report draws its own header.
    assert all(p.style.name != "Title" for p in document.paragraphs)

    heading = document.styles["Heading 1"]
    assert heading.font.name == "Arial"
    assert heading.font.color.rgb == RGBColor.from_string(PRIMARY_STRONG)
    # The template's theme font would override Arial, so it is dropped.
    rfonts = heading.element.rPr.find(qn("w:rFonts"))
    assert rfonts.get(qn("w:asciiTheme")) is None
    eyebrow = document.styles["WH Issue Eyebrow"]
    assert eyebrow.paragraph_format.page_break_before is (True if multi else None)

    # Status and label pills are shaded character styles.
    pill = document.styles[status_style("done")]
    assert pill.type == WD_STYLE_TYPE.CHARACTER
    assert pill.element.rPr.find(qn("w:shd")).get(qn("w:fill")) == STATUS_COLORS["done"][1]
    label = document.styles[label_style("bug")]
    assert label.element.rPr.find(qn("w:shd")).get(qn("w:fill")) == "F8DCDE"

    details, summary = document.tables
    assert details.cell(0, 0)._tc.tcPr.find(qn("w:shd")).get(qn("w:fill")) == SUNKEN
    assert summary.cell(0, 0)._tc.tcPr.find(qn("w:shd")).get(qn("w:fill")) == PRIMARY
    assert summary.autofit is False
    for table in (details, summary):
        assert len(table._tbl.tblPr.findall(qn("w:tblCellMar"))) == 1

    footer = document.sections[0].footer.paragraphs[0]
    assert "stale footer" not in footer.text
    assert footer.text.startswith("WikiHub issue report")
    fields = footer._p.xpath(".//w:instrText")
    assert [field.text.strip() for field in fields] == ["PAGE", "NUMPAGES"]


def test_style_issue_report_repaints_rather_than_stacking() -> None:
    """Running the pass over styles that already exist replaces their
    borders and shading instead of adding a second copy."""
    document = _load(_pandoc_like_docx())
    style_issue_report(document, multi=True, footer_text="Report")
    style_issue_report(document, multi=True, footer_text="Report")

    assert len(document.styles["WH Figure"].element.pPr.findall(qn("w:pBdr"))) == 1
    assert len(document.styles["WH Note"].element.pPr.findall(qn("w:shd"))) == 1
    assert len(document.styles[status_style("open")].element.rPr.findall(qn("w:shd"))) == 1
    summary = document.tables[1]
    assert len(summary.cell(0, 0)._tc.tcPr.findall(qn("w:shd"))) == 1
    assert len(summary._tbl.tblPr.findall(qn("w:tblCellMar"))) == 1


def _fake_pandoc(output: bytes) -> Any:
    """Stands in for the pandoc subprocess: "writes" ``output`` to its -o path."""

    async def create_subprocess_exec(*args: object, **_kwargs: object) -> Any:
        target = list(args)[list(args).index("-o") + 1]
        Path(str(target)).write_bytes(output)  # noqa: ASYNC240 - test double
        process = Mock(returncode=0)
        process.communicate = AsyncMock(return_value=(b"", b""))
        return process

    return patch("asyncio.create_subprocess_exec", side_effect=create_subprocess_exec)


async def test_issue_to_docx_styles_the_converted_document() -> None:
    with _fake_pandoc(_pandoc_like_docx()):
        content, filename = await issue_to_docx(_issue(), [], lang="vi")

    assert filename == "issue-save-fails.docx"
    document = _load(content)
    assert document.sections[0].footer.paragraphs[0].text.startswith("Báo cáo issue WikiHub")
    assert document.styles["WH Issue Eyebrow"].paragraph_format.page_break_before is None


async def test_issues_to_docx_builds_one_report_for_many_issues() -> None:
    issues = [(_issue(title="First"), []), (_issue(title="Second"), [])]
    with _fake_pandoc(_pandoc_like_docx()) as pandoc:
        content, filename = await issues_to_docx(issues, lang="en")

    assert filename.startswith("issues-2-") and filename.endswith(".docx")
    assert "title=Issues (2)" in pandoc.call_args.args
    document = _load(content)
    assert document.styles["WH Issue Eyebrow"].paragraph_format.page_break_before is True


async def test_issues_to_docx_keeps_the_single_issue_layout_for_one() -> None:
    with _fake_pandoc(_pandoc_like_docx()):
        _content, filename = await issues_to_docx([(_issue(), [])], lang="en")

    assert filename == "issue-save-fails.docx"


async def test_load_issue_images_skips_a_screenshot_that_is_gone() -> None:
    storage = Mock()
    storage.get = AsyncMock(side_effect=[b"png-bytes", FileNotFoundError("gone")])

    images = await load_issue_images(
        storage,
        [("a.png", "image/png", "key-a"), ("b.png", "image/png", "key-b")],
    )

    assert images == [("a.png", "image/png", b"png-bytes")]


def test_helpers_tolerate_missing_values() -> None:
    assert export_module._when(None) == ""
    # A run-properties element with no font entry at all is left alone.
    rpr = OxmlElement("w:rPr")
    export_style._clear_theme_fonts(rpr)
    assert rpr.find(qn("w:rFonts")) is None
