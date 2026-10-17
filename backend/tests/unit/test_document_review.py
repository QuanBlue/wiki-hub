"""Reviewing an imported page against its source file, and mending what differs."""

from __future__ import annotations

from pathlib import Path

import pymupdf
import pytest
from bs4 import BeautifulSoup
from docx import Document

from app.modules.document_import import review as rv
from app.modules.document_import.review import (
    PageBlock,
    Patch,
    SourceBlock,
    apply_patches,
    auto_fix,
    build_prompt,
    compare,
    parse_patches,
    read_page_blocks,
    read_source_blocks,
    review_import,
    tidy_visual,
)

SOURCE_HTML = """<html><head><title>x</title><style>p{}</style></head><body>
<h1>Quarterly report</h1>
<p>Revenue grew strongly this quarter across every region.</p>
<h2>Plans</h2>
<ul><li>Hire two engineers<ul><li>one backend</li></ul></li><li>Migrate the database</li></ul>
<pre>make deploy --prod</pre>
<blockquote>Quality is not an act.</blockquote>
<table><tr><th>Region</th><th>Growth</th></tr><tr><td>North</td><td>12%</td></tr></table>
<script>alert(1)</script>
</body></html>"""


def _docx(tmp_path: Path) -> Path:
    document = Document()
    document.add_heading("Quarterly report", level=1)
    document.add_paragraph("Revenue grew strongly this quarter across every region.")
    document.add_heading("Plans", level=2)
    document.add_paragraph("Hire two engineers", style="List Bullet")
    document.add_paragraph("Migrate the database", style="List Bullet")
    code = document.add_paragraph()
    code.add_run("make deploy --prod").font.name = "Courier New"
    document.add_paragraph("Quality is not an act.", style="Quote")
    table = document.add_table(rows=2, cols=2)
    for row, values in enumerate((("Region", "Growth"), ("North", "12%"))):
        for column, value in enumerate(values):
            table.cell(row, column).text = value
    document.add_paragraph("")  # an empty paragraph is not a block
    path = tmp_path / "report.docx"
    document.save(str(path))
    return path


def _kinds(blocks: list[SourceBlock]) -> list[tuple[str, str, int]]:
    return [(b.kind, b.text, b.level) for b in blocks]


# ------------------------------------------------------------------ reading the source


def test_a_word_file_is_read_as_it_was_written(tmp_path: Path) -> None:
    blocks = read_source_blocks(_docx(tmp_path), "docx")

    assert _kinds(blocks) == [
        ("heading", "Quarterly report", 1),
        ("paragraph", "Revenue grew strongly this quarter across every region.", 0),
        ("heading", "Plans", 2),
        ("list_item", "Hire two engineers", 0),
        ("list_item", "Migrate the database", 0),
        ("code", "make deploy --prod", 0),
        ("quote", "Quality is not an act.", 0),
        ("table", "Region Growth North 12%", 0),
    ]
    assert blocks[-1].rows == (("Region", "Growth"), ("North", "12%"))


def test_an_html_file_is_read_without_its_scripts_and_styles(tmp_path: Path) -> None:
    path = tmp_path / "page.html"
    path.write_text(SOURCE_HTML, encoding="utf-8")

    blocks = read_source_blocks(path, "html")

    assert _kinds(blocks) == [
        ("heading", "Quarterly report", 1),
        ("paragraph", "Revenue grew strongly this quarter across every region.", 0),
        ("heading", "Plans", 2),
        ("list_item", "Hire two engineers", 0),
        ("list_item", "one backend", 1),
        ("list_item", "Migrate the database", 0),
        ("code", "make deploy --prod", 0),
        ("quote", "Quality is not an act.", 0),
        ("table", "Region Growth North 12%", 0),
    ]


def test_a_pdf_is_read_by_blocks_and_what_repeats_on_each_page_is_furniture(
    tmp_path: Path,
) -> None:
    document = pymupdf.open()
    for number in range(1, 5):
        page = document.new_page()
        page.insert_text((72, 40), f"ACME CONFIDENTIAL {number}")
        page.insert_text(
            (72, 200), f"Body text of page {number} says something unique {number * 7}."
        )
    path = tmp_path / "doc.pdf"
    document.save(str(path))

    blocks = read_source_blocks(path, "pdf")

    furniture = [b for b in blocks if b.kind == "furniture"]
    body = [b for b in blocks if b.kind == "paragraph"]
    assert len(furniture) == 4 and all("ACME" in b.text for b in furniture)
    assert len(body) == 4 and "unique 7" in body[0].text
    assert read_source_blocks(path, "odt") == []


# ------------------------------------------------------------------------- the page


def test_the_page_is_read_by_its_top_level_blocks() -> None:
    page = read_page_blocks(
        "<h2>Title</h2><p>Text</p><ul><li>a<ul><li>b</li></ul></li></ul>"
        "<table><tr><th>x</th></tr></table><pre><code>c</code></pre>"
        "<blockquote><p>q</p></blockquote>"
        '<div data-type="callout"><p>n</p></div><hr>'
    )

    assert [(b.index, b.kind, b.text) for b in page] == [
        (0, "heading", "Title"),
        (1, "paragraph", "Text"),
        (2, "list", "a b"),
        (3, "table", "x"),
        (4, "code", "c"),
        (5, "quote", "q"),
        (6, "callout", "n"),
        (7, "other", ""),
    ]
    assert page[0].level == 2 and page[3].rows == (("x",),)


# ----------------------------------------------------------------------------- compare

GOOD = (
    "<h1>Quarterly report</h1><p>Revenue grew strongly this quarter across every region.</p>"
    "<h2>Plans</h2><ul><li>Hire two engineers</li><li>Migrate the database</li></ul>"
    "<pre><code>make deploy --prod</code></pre><blockquote><p>Quality is not an act.</p></blockquote>"
    "<table><tr><th>Region</th><th>Growth</th></tr><tr><td>North</td><td>12%</td></tr></table>"
)


def _source(tmp_path: Path) -> list[SourceBlock]:
    return read_source_blocks(_docx(tmp_path), "docx")


def test_a_faithful_page_has_no_differences(tmp_path: Path) -> None:
    result = compare(_source(tmp_path), read_page_blocks(GOOD))

    assert result.issues == [] and result.coverage == 1.0 and result.score == 0


def test_every_kind_of_difference_is_found(tmp_path: Path) -> None:
    broken = (
        "<p>Quarterly report</p>"  # a heading that became a paragraph
        "<h3>Plans</h3>"  # the wrong level
        "<p>Hire two engineers</p><p>Migrate the database</p>"  # a list gone flat
        "<p>make deploy --prod</p>"  # code as text
        "<table><tr><td>Region</td></tr></table>"  # a table with cells lost
        "<p>Confidential internal footer text</p>" * 3  # repeated furniture
    )  # the sentence about revenue and the quote are missing altogether

    result = compare(_source(tmp_path), read_page_blocks(broken))

    kinds = [issue.kind for issue in result.issues]
    assert kinds.count("missing") >= 2
    assert {"heading", "list", "code", "table", "extra"} <= set(kinds)
    assert result.missing == 2 and result.coverage < 1.0 and result.score > 10
    missing = [i.source.text for i in result.of("missing") if i.source]
    assert "Revenue grew strongly this quarter across every region." in missing
    assert "Quality is not an act." in missing


def test_blocks_in_another_order_are_one_difference(tmp_path: Path) -> None:
    source = [SourceBlock("paragraph", f"Paragraph number {n} of the document") for n in range(10)]
    page = read_page_blocks(
        "".join(
            f"<p>Paragraph number {n} of the document</p>" for n in (5, 6, 7, 8, 9, 0, 1, 2, 3, 4)
        )
    )
    scrambled = read_page_blocks(
        "".join(f"<p>Paragraph number {n} of the document</p>" for n in reversed(range(10)))
    )

    assert not compare(source, page).of("order")
    result = compare(source, scrambled)
    assert [i.kind for i in result.of("order")] == ["order"]


def test_a_table_of_another_shape_is_a_difference(tmp_path: Path) -> None:
    source = [SourceBlock("table", "a b c d", rows=(("a", "b"), ("c", "d")))]

    wrong = compare(
        source, read_page_blocks("<table><tr><td>a b</td></tr><tr><td>c d</td></tr></table>")
    )
    right = compare(
        source,
        read_page_blocks(
            "<table><tr><td>a</td><td>b</td></tr><tr><td>c</td><td>d</td></tr></table>"
        ),
    )

    assert [i.kind for i in wrong.issues] == ["table"] and "2 rows by 2 columns" in wrong.issues[
        0
    ].detail
    assert right.issues == []


def test_small_changes_of_wording_are_not_a_loss() -> None:
    source = [SourceBlock("paragraph", "The “quick” brown fox, jumps over the lazy dog today")]

    result = compare(
        source, read_page_blocks("<p>The quick brown fox jumps over the lazy dog today.</p>")
    )

    assert result.issues == [] and result.coverage == 1.0


# ----------------------------------------------------------------------------- auto_fix


def test_what_is_certain_is_mended_without_a_model(tmp_path: Path) -> None:
    source = _source(tmp_path)
    broken = (
        "<p>Quarterly report</p><h3>Plans</h3>"
        "<ul><li>Hire two engineers</li><li>Migrate the database</li></ul>"
        "<pre><code>make deploy --prod</code></pre>"
        "<table><tr><th>Region</th><th>Growth</th></tr><tr><td>North</td><td>12%</td></tr></table>"
        + "<p>Confidential internal footer text</p>"
        * 3
    )
    before = compare(source, read_page_blocks(broken))

    mended = auto_fix(broken, before)

    after = compare(source, read_page_blocks(mended.html))
    assert after.score < before.score and after.missing == 0
    assert mended.fixes >= 6
    assert "<h1>Quarterly report</h1>" in mended.html and "<h2>Plans</h2>" in mended.html
    assert "Confidential internal footer" not in mended.html
    # The lost sentences are back, after the block that came before them in the source.
    assert mended.html.index("Revenue grew") > mended.html.index("Quarterly report")
    assert mended.html.index("Revenue grew") < mended.html.index("Plans")
    assert any("heading level" in note for note in mended.notes)
    assert any("missing" in note for note in mended.notes)
    assert any("repeated" in note for note in mended.notes)


def test_lost_blocks_come_back_in_their_own_form() -> None:
    source = [
        SourceBlock("heading", "A heading here", 3),
        SourceBlock("list_item", "first item of the list"),
        SourceBlock("code", "print('hello world')"),
        SourceBlock("quote", "a quoted sentence here"),
        SourceBlock("table", "x y", rows=(("x", "y"),)),
        SourceBlock("paragraph", "last paragraph of text"),
    ]
    before = compare(source, read_page_blocks("<p>Something else entirely different</p>"))

    html = auto_fix("<p>Something else entirely different</p>", before).html

    for fragment in (
        "<h3>A heading here</h3>",
        "<ul><li>first item of the list</li></ul>",
        "<pre><code>print('hello world')</code></pre>",
        "<blockquote><p>a quoted sentence here</p></blockquote>",
        "<th><p>x</p></th>",
        "<p>last paragraph of text</p>",
    ):
        assert fragment in html
    assert compare(source, read_page_blocks(html)).missing == 0


def test_text_the_source_lacks_is_only_removed_when_it_repeats() -> None:
    source = [SourceBlock("paragraph", "the only thing the document says")]
    page = "<p>the only thing the document says</p><p>one lonely stray line of text</p>"
    before = compare(source, read_page_blocks(page))

    mended = auto_fix(page, before)

    # A single stray line may be the author's: it is left for a person (or the model) to judge.
    assert "lonely stray" in mended.html and mended.fixes == 0


# ------------------------------------------------------------------------------ patches


def test_patches_are_read_from_a_reply_and_what_is_malformed_is_dropped() -> None:
    reply = """Here you go:
    [{"op":"replace","block":3,"markdown":"| a | b |\\n|---|---|\\n| 1 | 2 |"},
     {"op":"delete","block":5},
     {"op":"heading","block":1,"level":2},
     {"op":"heading","block":1,"level":9},
     {"op":"replace","block":2,"markdown":"  "},
     {"op":"explode","block":2},
     {"op":"delete","block":"x"},
     "junk"]"""

    patches = parse_patches(reply)

    assert patches == [
        Patch("replace", 3, "| a | b |\n|---|---|\n| 1 | 2 |"),
        Patch("delete", 5),
        Patch("heading", 1, level=2),
    ]
    for reply in ("", "no json", "[not json]", '{"op":"delete"}', "[]"):
        assert parse_patches(reply) == []


PAGE = "<h1>T</h1><p>one</p><p>two</p><p>three</p><p>four</p>"


def test_patches_replace_insert_delete_and_retitle() -> None:
    patches = [
        Patch("replace", 1, "| a | b |\n|---|---|\n| 1 | 2 |"),
        Patch("insert_after", 2, "- x\n- y"),
        Patch("delete", 3),
        Patch("heading", 4, level=3),
        Patch("delete", 99),  # no such block
        Patch("heading", 0, level=2),
    ]

    html, applied = apply_patches(PAGE, patches)

    assert applied == 5
    assert "<table>" in html and "<th>a</th>" in html
    assert "<ul>" in html and "<li>x</li>" in html and "three" not in html
    assert "<h3>four</h3>" in html and "<h2>T</h2>" in html
    assert html.index("<table>") < html.index("two") < html.index("<ul>")


def test_a_callout_written_as_markdown_becomes_the_editors_callout() -> None:
    html, applied = apply_patches(
        "<p>x</p>", [Patch("replace", 0, "> [!WARNING] Back up first.\n> Then continue.")]
    )

    assert applied == 1
    assert 'data-type="callout"' in html and 'data-callout-type="warning"' in html
    assert "Back up first." in html and "[!WARNING]" not in html

    plain, _ = apply_patches("<p>x</p>", [Patch("replace", 0, "> just a quote")])
    assert "<blockquote>" in plain and "callout" not in plain


def test_a_task_list_and_code_block_use_the_editors_shapes() -> None:
    html, _ = apply_patches(
        "<p>x</p>", [Patch("replace", 0, "- [ ] todo\n- [x] done\n\n```python\nprint(1)\n```")]
    )

    assert 'data-type="taskList"' in html and 'data-checked="true"' in html
    assert "<pre>" in html and "print(1)" in html


def test_a_patch_never_brings_script_into_the_page() -> None:
    html, _ = apply_patches("<p>x</p>", [Patch("replace", 0, "<script>alert(1)</script>hello")])

    assert "<script" not in html and "hello" in html


# --------------------------------------------------------------------------- the prompt


def test_the_prompt_shows_each_issue_with_the_original_and_the_blocks_around_it(
    tmp_path: Path,
) -> None:
    source = _source(tmp_path)
    page = read_page_blocks(
        "<p>a</p><p>b</p><p>c</p><p>d</p><p>e</p><p>f</p><p>g</p>"
        "<table><tr><td>Region Growth</td></tr></table>"
    )
    issues = compare(source, page).of("table")

    prompt = build_prompt(issues, page)

    assert "<issues>" in prompt and "table:" in prompt
    assert "ORIGINAL (table): Region | Growth\nNorth | 12%" in prompt
    assert "PAGE block: [7]" in prompt
    assert "[7] table:" in prompt and "[0] paragraph" not in prompt  # only the neighbourhood
    assert "never add a word" in rv.REVIEW_RULES and "[!NOTE]" in rv.REVIEW_RULES


# ----------------------------------------------------------------------- the whole review


class Model:
    def __init__(self, *replies: str | None) -> None:
        self.replies = list(replies)
        self.asked: list[str] = []

    async def __call__(self, system: str, user: str) -> str | None:
        self.asked.append(user)
        return self.replies.pop(0) if self.replies else "[]"


# A table that fell apart into one cell: a form the conversion left, not a paragraph to reshape.
BROKEN_TABLE = (
    "<h1>Quarterly report</h1><p>Revenue grew strongly this quarter across every region.</p>"
    "<h2>Plans</h2><ul><li>Hire two engineers</li><li>Migrate the database</li></ul>"
    "<pre><code>make deploy --prod</code></pre><blockquote><p>Quality is not an act.</p></blockquote>"
    "<table><tr><td>Region Growth North 12%</td></tr></table>"
)
# The same table written out as a paragraph.
FLAT_TABLE = BROKEN_TABLE.replace(
    "<table><tr><td>Region Growth North 12%</td></tr></table>", "<p>Region Growth North 12%</p>"
)


async def test_a_clean_page_costs_nothing(tmp_path: Path) -> None:
    model = Model()

    review = await review_import(GOOD, _docx(tmp_path), "docx", ask=model)

    assert review.status == "clean" and review.html == GOOD and model.asked == []


async def test_what_is_certain_is_mended_and_no_model_is_asked_for_it(tmp_path: Path) -> None:
    model = Model()
    broken = GOOD.replace("<h2>Plans</h2>", "<p>Plans</p>")

    review = await review_import(broken, _docx(tmp_path), "docx", ask=model)

    assert review.status == "fixed" and "<h2>Plans</h2>" in review.html
    assert review.fixes == 1 and model.asked == [] and review.issues_after == 0


async def test_a_table_written_out_as_a_paragraph_needs_no_model(tmp_path: Path) -> None:
    model = Model()

    review = await review_import(FLAT_TABLE, _docx(tmp_path), "docx", ask=model)

    assert review.status == "fixed" and model.asked == [] and review.ai_calls == 0
    assert "<th><p>Region</p></th>" in review.html


async def test_the_model_rebuilds_what_cannot_be_mended_without_it(tmp_path: Path) -> None:
    table = '[{"op":"replace","block":6,"markdown":"| Region | Growth |\\n|---|---|\\n| North | 12% |"}]'
    model = Model(table)

    review = await review_import(BROKEN_TABLE, _docx(tmp_path), "docx", ask=model)

    assert review.status == "fixed" and review.ai_calls == 1
    assert "<th>Region</th>" in review.html and "<td>12%</td>" in review.html
    assert "Region Growth North 12%" not in review.html
    assert "ORIGINAL (table)" in model.asked[0]
    assert any("AI assistant mended 1 block" in note for note in review.notes)
    assert review.issues_before >= 1 and review.issues_after == 0


async def test_a_change_that_loses_text_is_thrown_away(tmp_path: Path) -> None:
    # The model "tidies" by deleting the table and a sentence: the page matches the source less.
    harmful = '[{"op":"delete","block":6},{"op":"delete","block":1}]'
    model = Model(harmful)

    review = await review_import(BROKEN_TABLE, _docx(tmp_path), "docx", ask=model)

    assert review.html == BROKEN_TABLE and review.status == "kept"
    assert any("did not match better" in note for note in review.notes)


async def test_without_a_model_only_the_certain_is_done_and_the_rest_is_reported(
    tmp_path: Path,
) -> None:
    review = await review_import(
        BROKEN_TABLE.replace("<h2>Plans</h2>", "<p>Plans</p>"), _docx(tmp_path), "docx"
    )

    assert review.status == "partial" and "<h2>Plans</h2>" in review.html
    assert review.issues_after >= 1 and any("left" in note for note in review.notes)


async def test_a_model_that_is_not_available_or_answers_nonsense_changes_nothing(
    tmp_path: Path,
) -> None:
    unavailable = await review_import(BROKEN_TABLE, _docx(tmp_path), "docx", ask=Model(None))
    assert unavailable.html == BROKEN_TABLE and unavailable.status == "kept"
    assert any("not available" in note for note in unavailable.notes)

    nonsense = await review_import(
        BROKEN_TABLE, _docx(tmp_path), "docx", ask=Model("I cannot", "[]")
    )
    assert nonsense.html == BROKEN_TABLE and nonsense.ai_calls >= 1

    bad_block = await review_import(
        BROKEN_TABLE, _docx(tmp_path), "docx", ask=Model('[{"op":"delete","block":99}]')
    )
    assert bad_block.html == BROKEN_TABLE


async def test_the_model_is_asked_at_most_a_few_times_for_one_document(tmp_path: Path) -> None:
    source = [f"<p>Paragraph {n} of a long document with several words</p>" for n in range(60)]
    path = tmp_path / "long.html"
    path.write_text("<html><body>" + "".join(source) + "</body></html>", encoding="utf-8")
    page = "<p>unrelated words only appear here</p>" * 2
    model = Model()

    review = await review_import(page, path, "html", ask=model)

    assert review.ai_calls <= rv.MAX_AI_CALLS and len(model.asked) <= rv.MAX_AI_CALLS


async def test_what_cannot_be_read_again_or_is_not_reviewable_is_skipped(tmp_path: Path) -> None:
    other = await review_import("<p>x</p>", tmp_path / "a.odt", "odt")
    assert other.status == "skipped" and other.html == "<p>x</p>"

    broken = tmp_path / "broken.docx"
    broken.write_bytes(b"this is not a zip")
    unreadable = await review_import("<p>x</p>", broken, "docx")
    assert unreadable.status == "skipped" and "could not be read" in unreadable.notes[0]

    empty = tmp_path / "empty.html"
    empty.write_text("<html><body></body></html>", encoding="utf-8")
    assert (await review_import("<p>x</p>", empty, "html")).status == "skipped"


async def test_a_pdf_page_with_repeated_headers_and_a_lost_paragraph_is_mended(
    tmp_path: Path,
) -> None:
    document = pymupdf.open()
    for number in range(1, 5):
        page = document.new_page()
        page.insert_text((72, 40), f"ACME CONFIDENTIAL {number}")
        page.insert_text(
            (72, 200), f"Body text of page {number} says something unique {number * 7}."
        )
    path = tmp_path / "doc.pdf"
    document.save(str(path))
    converted = (
        "<p>ACME CONFIDENTIAL</p><p>Body text of page 1 says something unique 7.</p>"
        "<p>Body text of page 2 says something unique 14.</p>"
        "<p>Body text of page 4 says something unique 28.</p>"
        "<p>ACME CONFIDENTIAL</p><p>ACME CONFIDENTIAL</p>"
    )

    review = await review_import(converted, path, "pdf")

    assert "ACME CONFIDENTIAL" not in review.html
    assert "unique 21" in review.html  # page 3's paragraph is back
    assert review.status in {"fixed", "partial"} and review.fixes >= 2


def test_a_review_is_recorded_in_a_small_form_for_the_job() -> None:
    review = rv.Review("fixed", "<p>x</p>", 5, 1, 4, 2, [f"note {n}" for n in range(12)])

    record = review.as_record()

    assert record["status"] == "fixed" and record["fixes"] == 4 and len(record["notes"]) == 8
    assert set(record) == {
        "status",
        "issues_before",
        "issues_after",
        "fixes",
        "ai_calls",
        "notes",
        "log",
    }
    assert PageBlock(0, "other", "").html == ""


@pytest.mark.parametrize("text", ["Đường Nguyễn Trãi", "ĐƯỜNG NGUYỄN TRÃI", "duong nguyen trai"])
def test_comparison_ignores_case_and_accents(text: str) -> None:
    assert rv.norm(text) == "duong nguyen trai"


def test_forms_the_conversion_dropped_are_given_back_when_the_text_is_the_same() -> None:
    source = [
        SourceBlock("code", "make deploy --prod"),
        SourceBlock("quote", "Quality is not an act."),
        SourceBlock("list_item", "first item of the list"),
        SourceBlock("list_item", "second item of the list"),
        SourceBlock("paragraph", "a paragraph in between the two lists"),
        SourceBlock("list_item", "item of another list"),
        SourceBlock("table", "x y z w", rows=(("x", "y"), ("z", "w"))),
    ]
    flat = (
        "<p>make deploy --prod</p><p>Quality is not an act.</p>"
        "<p>first item of the list</p><p>second item of the list</p>"
        "<p>a paragraph in between the two lists</p><p>item of another list</p>"
        "<p>x y z w</p>"
    )
    before = compare(source, read_page_blocks(flat))

    mended = auto_fix(flat, before)

    assert mended.html == (
        "<pre><code>make deploy --prod</code></pre>"
        "<blockquote><p>Quality is not an act.</p></blockquote>"
        "<ul><li>first item of the list</li><li>second item of the list</li></ul>"
        "<p>a paragraph in between the two lists</p>"
        "<ul><li>item of another list</li></ul>"
        "<table><tbody><tr><th><p>x</p></th><th><p>y</p></th></tr>"
        "<tr><td><p>z</p></td><td><p>w</p></td></tr></tbody></table>"
    )
    assert compare(source, read_page_blocks(mended.html)).issues == []
    # Seven blocks, one of which (the paragraph between the lists) was right already.
    assert mended.fixes == 6 and any("their form again" in note for note in mended.notes)


def test_a_form_is_only_given_back_where_the_page_says_the_same_as_the_source() -> None:
    source = [SourceBlock("code", "alpha beta gamma delta epsilon zeta eta")]
    # Nearly the same words, one short: found, but not certain enough to be reshaped.
    page = "<p>alpha beta gamma delta epsilon zeta</p>"
    before = compare(source, read_page_blocks(page))

    assert [i.kind for i in before.issues] == ["code"]
    mended = auto_fix(page, before)
    assert mended.fixes == 0 and mended.html == page


def test_a_page_number_that_changes_on_every_page_is_still_a_repeated_footer() -> None:
    source = [
        SourceBlock("paragraph", f"Body paragraph number {n} with some words") for n in range(4)
    ]
    page = "".join(
        f"<p>Body paragraph number {n} with some words</p><p>Page {n + 1} of 4</p>"
        for n in range(4)
    )
    before = compare(source, read_page_blocks(page))

    mended = auto_fix(page, before)

    assert "Page " not in mended.html and mended.html.count("Body paragraph") == 4


def test_the_footer_of_a_pdf_goes_even_when_its_words_are_used_by_the_body() -> None:
    # "Page 3 of 4": every one of its words is in the body text too.
    source = [
        SourceBlock("paragraph", "The paragraph of page 1 says one thing of 4"),
        SourceBlock("furniture", "Page 1 of 4"),
        SourceBlock("paragraph", "The paragraph of page 2 says another thing of 4"),
        SourceBlock("furniture", "Page 2 of 4"),
        SourceBlock("paragraph", "The paragraph of page 3 says a third thing of 4"),
        SourceBlock("furniture", "Page 3 of 4"),
    ]
    page = (
        "<p>The paragraph of page 1 says one thing of 4</p><p>Page 1 of 4</p>"
        "<p>The paragraph of page 2 says another thing of 4</p><p>Page 2 of 4</p>"
        "<p>The paragraph of page 3 says a third thing of 4</p><p>Page 3 of 4</p>"
    )
    before = compare(source, read_page_blocks(page))

    mended = auto_fix(page, before)

    assert [i.kind for i in before.issues] == ["extra"] * 3
    assert "Page " not in mended.html.replace("page 1", "").replace("page 2", "").replace(
        "page 3", ""
    )
    assert mended.html.count("<p>") == 3
    # A body line that merely resembles a footer is not touched when the source has it as text.
    kept = compare(
        [SourceBlock("paragraph", "Page 1 of 4")], read_page_blocks("<p>Page 1 of 4</p>")
    )
    assert kept.issues == []


# ------------------------------------------------------------- the edges of reading and mending


def test_a_word_title_and_a_merged_cell_are_read_as_they_are(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("The big title", style="Title")
    table = document.add_table(rows=2, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "merged heading"
    table.cell(0, 2).text = "side"
    for column, value in enumerate(("a", "b", "c")):
        table.cell(1, column).text = value
    path = tmp_path / "merged.docx"
    document.save(str(path))

    blocks = read_source_blocks(path, "docx")

    assert (blocks[0].kind, blocks[0].level) == ("heading", 1)
    # The merged cell is one cell of its row, not one per column it spans.
    assert blocks[1].rows == (("merged heading", "side"), ("a", "b", "c"))


def test_loose_text_and_wrapped_blocks_of_an_html_file_are_read(tmp_path: Path) -> None:
    path = tmp_path / "wrapped.html"
    path.write_text(
        "<html><body>loose text before<div><section><p>inside wrappers</p></section></div>"
        "<table><tr><td></td></tr></table></body></html>",
        encoding="utf-8",
    )

    blocks = read_source_blocks(path, "html")

    # An empty table is nothing; the loose words and the wrapped paragraph are blocks.
    assert [(b.kind, b.text) for b in blocks] == [
        ("paragraph", "loose text before"),
        ("paragraph", "inside wrappers"),
    ]


def test_a_pdf_is_read_up_to_a_page_limit_and_without_its_pictures(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    document = pymupdf.open()
    pixels = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 20, 20), False)
    pixels.clear_with(200)
    for number in range(4):
        page = document.new_page()
        page.insert_image(pymupdf.Rect(72, 300, 120, 348), pixmap=pixels)
        page.insert_text((72, 200), f"Text of page {number + 1} only")
    path = tmp_path / "pictures.pdf"
    document.save(str(path))
    monkeypatch.setattr(rv, "_PDF_PAGES", 2)

    blocks = read_source_blocks(path, "pdf")

    assert [b.text for b in blocks] == ["Text of page 1 only", "Text of page 2 only"]


def test_short_pieces_of_text_and_pages_with_a_rule_are_not_compared() -> None:
    source = [SourceBlock("paragraph", "1"), SourceBlock("paragraph", "a real sentence of text")]

    result = compare(source, read_page_blocks("<hr><p>a real sentence of text</p><hr>"))

    assert result.issues == [] and result.coverage == 1.0


def test_lost_text_goes_into_an_empty_page_and_ahead_of_everything_when_nothing_came_first() -> (
    None
):
    source = [SourceBlock("paragraph", "the very first words of the document")]

    empty = auto_fix("", compare(source, read_page_blocks("")))
    first = auto_fix("<p>other</p>", compare(source, read_page_blocks("<p>other</p>")))

    assert empty.html == "<p>the very first words of the document</p>"
    assert first.html.startswith("<p>the very first words of the document</p>")


def test_a_patch_that_cannot_apply_is_ignored() -> None:
    page = "<ul><li>a</li></ul><p>b</p>"

    html, applied = apply_patches(
        page,
        [
            Patch("heading", 0, level=2),  # a list is not a heading
        ],
    )

    assert applied == 0 and html == page


def test_the_prompt_says_where_lost_text_belongs() -> None:
    page = read_page_blocks("<p>alpha one</p><p>beta two</p><p>gamma three</p>")
    first = compare([SourceBlock("paragraph", "lost at the very start")], page).issues
    later = compare(
        [
            SourceBlock("paragraph", "alpha one"),
            SourceBlock("paragraph", "lost in the middle of it"),
        ],
        page,
    ).issues

    assert "belongs after PAGE block [-1]" in build_prompt(first, page)
    assert "belongs after PAGE block [0]" in build_prompt(later, page)
    assert "[0] paragraph: alpha one" in build_prompt(first, page)


async def test_the_review_says_when_more_was_left_than_it_could_look_at(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(rv, "MAX_AI_CALLS", 1)
    monkeypatch.setattr(rv, "MAX_ISSUES_PER_CALL", 1)
    path = tmp_path / "many.docx"
    document = Document()
    for number in range(3):
        document.add_paragraph(f"Item number {number} of the list", style="List Bullet")
    document.save(str(path))
    model = Model("[]", "[]", "[]")

    review = await review_import("<p>unrelated text only here</p>" * 2, path, "docx", ask=model)

    assert review.ai_calls == 1 and len(model.asked) == 1
    assert any("more differences were left" in note for note in review.notes)


# ---------------------------------------------------------------- the log, and finding the block


async def test_the_review_tells_each_step_as_it_goes(tmp_path: Path) -> None:
    table = '[{"op":"replace","block":6,"markdown":"| Region | Growth |\\n|---|---|\\n| North | 12% |"}]'
    seen: list[list[str]] = []

    async def report(lines: list[str]) -> None:
        seen.append(list(lines))

    review = await review_import(
        BROKEN_TABLE, _docx(tmp_path), "docx", ask=Model(table), report=report
    )

    text = "\n".join(review.log)
    assert "Reading the original file again" in text
    assert "Compared" in text and "1 table" in text
    assert "Asking the AI assistant about 1 difference(s)" in text
    assert "proposed 1 change(s); 1 applied and kept" in text
    assert review.log[-1].startswith("Done:") and "now matches the original" in review.log[-1]
    # Told after every step, each time with the log so far.
    assert len(seen) == len(review.log) and [len(lines) for lines in seen] == list(
        range(1, len(review.log) + 1)
    )
    assert review.as_record()["log"] == review.log


async def test_the_log_says_why_a_change_was_dropped_or_the_ai_was_not_there(
    tmp_path: Path,
) -> None:
    harmful = '[{"op":"delete","block":6},{"op":"delete","block":1}]'

    dropped = await review_import(BROKEN_TABLE, _docx(tmp_path), "docx", ask=Model(harmful))
    without = await review_import(BROKEN_TABLE, _docx(tmp_path), "docx")
    nothing = await review_import(BROKEN_TABLE, _docx(tmp_path), "docx", ask=Model("[]"))

    assert any("were dropped" in line for line in dropped.log)
    assert dropped.log[-1] == "Done: the page is left as it was converted."
    assert any("AI assistant is not set up" in line for line in without.log)
    assert any("proposed no change" in line for line in nothing.log)


async def test_a_skipped_review_still_has_a_log(tmp_path: Path) -> None:
    review = await review_import("<p>x</p>", tmp_path / "a.odt", "odt")

    assert review.status == "skipped" and review.log == ["This file type is not reviewed."]


def test_a_short_item_is_matched_to_the_list_not_to_a_table_that_has_the_word() -> None:
    source = [SourceBlock("list_item", "Giám sát")]
    page = read_page_blocks(
        "<table><tr><td>Giám sát</td><td>Quản lý</td></tr></table><ul><li>Giám sát</li></ul>"
    )

    assert compare(source, page).issues == []


def test_a_heading_is_matched_to_the_heading_not_to_its_line_in_the_contents() -> None:
    source = [SourceBlock("heading", "Kết nối ngoài", level=2)]
    page = read_page_blocks("<h1>Kết nối ngoài 131</h1><p>text</p><h2>Kết nối ngoài</h2>")

    assert compare(source, page).issues == []


def test_a_table_that_kept_every_cell_is_not_broken_for_having_more_rows() -> None:
    source = [SourceBlock("table", "a b c d", rows=(("a", "b"), ("c", "d")))]
    page = read_page_blocks(
        "<table><tr><td>a</td><td>b</td><td></td></tr><tr><td>1</td><td></td><td></td></tr>"
        "<tr><td>c</td><td>d</td><td></td></tr></table>"
    )

    assert compare(source, page).issues == []
    # Losing a column, though, is.
    narrow = read_page_blocks("<table><tr><td>a b</td></tr><tr><td>c d</td></tr></table>")
    assert [i.kind for i in compare(source, narrow).issues] == ["table"]


# ------------------------------------------------------------------- how the page looks


def test_an_indented_line_in_a_table_cell_is_not_left_as_a_quote() -> None:
    page = (
        '<table><tr><th style="background-color:#2f5496;color:#e7e6e6">'
        "<blockquote><p><strong>Phân hệ</strong></p></blockquote></th>"
        "<td><p>text</p></td></tr></table><blockquote><p>a real quote</p></blockquote>"
    )

    result = tidy_visual(page)

    assert "<th" in result.html and "<strong>Phân hệ</strong>" in result.html
    # The quote outside any cell is a quotation and stays.
    assert result.html.count("<blockquote>") == 1 and "a real quote" in result.html
    assert result.fixes == 1 and "out of a quote" in result.notes[0]


def test_text_on_a_coloured_cell_gets_a_colour_that_reads() -> None:
    page = (
        '<table><tr><td style="background-color:#2f5496">no colour</td>'
        '<td style="background-color:#2f5496;color:#3a5ba0">too close</td>'
        '<td style="background-color:#ffff00;color:#ffffcc">pale on yellow</td>'
        '<td style="background-color:#2f5496;color:#e7e6e6">fine</td>'
        '<td style="background-color:#deeaf6">light, no colour</td></tr></table>'
    )

    result = tidy_visual(page)

    cells = BeautifulSoup(result.html, "html.parser").find_all("td")
    assert "color:#ffffff" in cells[0]["style"] and "color:#ffffff" in cells[1]["style"]
    assert "color:#000000" in cells[2]["style"]
    assert cells[3]["style"] == "background-color:#2f5496;color:#e7e6e6"
    assert cells[4]["style"] == "background-color:#deeaf6"
    assert result.fixes == 1 and "readable" in result.notes[0]


def test_a_forced_black_is_dropped_only_where_nothing_is_filled_behind_it() -> None:
    page = (
        '<p style="color:#000000;text-align:justify">free</p>'
        '<table><tr><td style="background-color:#deeaf6">'
        '<p style="color:#000000">on a fill</p></td></tr></table>'
    )

    result = tidy_visual(page)

    soup = BeautifulSoup(result.html, "html.parser")
    free, filled = soup.find_all("p")
    assert free["style"] == "text-align:justify"
    assert filled["style"] == "color:#000000"
    assert result.fixes == 1 and "forced black" in result.notes[0]


def test_a_page_that_looks_fine_is_returned_as_it_is() -> None:
    page = '<p style="text-align:center">fine</p><table><tr><td>x</td></tr></table>'

    result = tidy_visual(page)

    assert result.html == page and result.fixes == 0 and result.notes == []


def test_colours_that_are_not_hex_are_left_alone() -> None:
    page = '<table><tr><td style="background-color:rgb(1,2,3)">x</td></tr></table>'

    assert tidy_visual(page).fixes == 0


async def test_a_page_whose_text_matches_still_has_its_look_tidied(tmp_path: Path) -> None:
    model = Model()
    page = GOOD.replace(
        "<table>", '<table><thead></thead>'
    ).replace("<th>Region</th>", '<th style="background-color:#2f5496">Region</th>')

    review = await review_import(page, _docx(tmp_path), "docx", ask=model)

    assert review.status == "fixed" and model.asked == []
    assert "color:#ffffff" in review.html
    assert any("Fixed without AI: made the text of 1 coloured cell readable" in line for line in review.log)
    assert review.notes[-1] == "the page matches the original (its look was tidied)"


def test_a_style_with_nothing_left_is_removed_and_short_hex_colours_are_read() -> None:
    page = (
        '<p style="color:#000">only a black</p>'
        '<table><tr><td style="background-color:#000">on black</td></tr></table>'
    )

    result = tidy_visual(page)

    soup = BeautifulSoup(result.html, "html.parser")
    assert soup.find("p").get("style") is None
    assert "color:#ffffff" in soup.find("td")["style"]


# ------------------------------------------------------------------- how deep a list item sits


def _depths(path: Path, fmt: str) -> list[int]:
    return [b.level for b in read_source_blocks(path, fmt) if b.kind == "list_item"]


def test_a_word_list_is_nested_by_how_far_in_it_sits(tmp_path: Path) -> None:
    from docx.shared import Inches

    document = Document()
    for text, left in (("a", 0.25), ("b", 0.5), ("c", 0.5), ("d", 0.25), ("e", 0.75)):
        paragraph = document.add_paragraph(text, style="List Bullet")
        paragraph.paragraph_format.left_indent = Inches(left)
    path = tmp_path / "indent.docx"
    document.save(str(path))

    assert _depths(path, "docx") == [0, 1, 1, 0, 1]


def test_a_word_list_is_nested_by_its_level_and_its_style(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("top", style="List Bullet")
    deep = document.add_paragraph("deep", style="List Bullet")
    number = deep._p.get_or_add_pPr().get_or_add_numPr()
    number.get_or_add_ilvl().val = 1
    number.get_or_add_numId().val = document.styles["List Bullet"].element.pPr.numPr.numId.val
    document.add_paragraph("by style", style="List Bullet 2")
    path = tmp_path / "levels.docx"
    document.save(str(path))

    assert _depths(path, "docx") == [0, 1, 1]


def test_a_list_broken_by_a_paragraph_starts_again() -> None:
    blocks = [
        SourceBlock("list_item", "a"),
        SourceBlock("list_item", "b"),
        SourceBlock("paragraph", "between"),
        SourceBlock("list_item", "c"),
    ]

    result = rv._with_depths(blocks, [(0, 720), (1, 1440), (3, 1440)])

    assert [b.level for b in result if b.kind == "list_item"] == [0, 1, 0]


NESTED_SOURCE = (
    "<html><body><ul><li>Server policy<ul><li>everyone has a password</li>"
    "<li>only root has uid 0</li></ul></li><li>Application policy</li></ul></body></html>"
)


async def test_a_list_the_conversion_made_flat_is_nested_again_without_a_model(
    tmp_path: Path,
) -> None:
    path = tmp_path / "nested.html"
    path.write_text(NESTED_SOURCE, encoding="utf-8")
    flat = (
        "<ul><li><strong>Server policy</strong></li><li>everyone has a password</li>"
        "<li>only root has uid 0</li><li>Application policy</li></ul>"
    )
    model = Model()

    review = await review_import(flat, path, "html", ask=model)

    assert review.status == "fixed" and model.asked == [] and review.issues_after == 0
    soup = BeautifulSoup(review.html, "html.parser")
    top = soup.find("ul")
    assert [li.find(string=True, recursive=False) for li in top.find_all("li", recursive=False)][
        1
    ] == "Application policy"
    nested = top.find("li").find("ul")
    assert [li.text for li in nested.find_all("li")] == [
        "everyone has a password",
        "only root has uid 0",
    ]
    # The formatting of the item that kept its place is kept.
    assert top.find("li").find("strong").text == "Server policy"
    assert any("nesting back" in note for note in review.notes)
    assert any("list nesting" in line for line in review.log)


def test_lists_that_follow_one_another_are_one_list_to_nest() -> None:
    source = [SourceBlock("list_item", "alpha", 0), SourceBlock("list_item", "beta", 1)]
    page = read_page_blocks("<ol><li>alpha</li></ol><ol><li>beta</li></ol>")

    issues = compare(source, page).issues
    mended = auto_fix("<ol><li>alpha</li></ol><ol><li>beta</li></ol>", compare(source, page))

    assert [i.kind for i in issues] == ["indent"] and issues[0].plan.blocks == [0, 1]
    assert mended.html == "<ol><li>alpha<ol><li>beta</li></ol></li></ol>"


def test_a_deeper_step_than_one_level_is_kept_to_one_and_other_lists_are_left() -> None:
    source = [
        SourceBlock("list_item", "alpha", 0),
        SourceBlock("list_item", "beta", 2),
        SourceBlock("list_item", "other", 0),
    ]
    page = read_page_blocks(
        "<ul><li>alpha</li><li>beta</li></ul><p>x</p><ul><li>unrelated</li></ul>"
    )

    [issue] = [i for i in compare(source, page).issues if i.kind == "indent"]

    assert issue.plan.desired == [0, 1] and issue.plan.blocks == [0]
    # A list whose items are not the source's in order is not touched at all.
    swapped = compare(source, read_page_blocks("<ul><li>beta</li><li>alpha</li></ul>"))
    assert [i.kind for i in swapped.issues if i.kind == "indent"] == []


def test_a_list_already_nested_as_the_source_has_no_issue() -> None:
    source = [SourceBlock("list_item", "alpha", 0), SourceBlock("list_item", "beta", 1)]
    page = read_page_blocks("<ul><li>alpha<ul><li>beta</li></ul></li></ul>")

    assert compare(source, page).issues == []


def test_a_list_that_vanished_while_mending_is_not_rebuilt() -> None:
    source = [SourceBlock("list_item", "alpha", 0), SourceBlock("list_item", "beta", 1)]
    html = "<ul><li>alpha</li><li>beta</li></ul>"
    comparison = compare(source, read_page_blocks(html))
    comparison.issues[0].plan.blocks = [5]  # a block that is no longer there

    assert auto_fix(html, comparison).html == html


def test_a_list_style_with_an_indent_of_its_own_and_one_without(tmp_path: Path) -> None:
    document = Document()
    for text, style in (("flush", "List"), ("further in", "List Paragraph")):
        document.add_paragraph(text, style=style)
    path = tmp_path / "styles.docx"
    document.save(str(path))

    assert _depths(path, "docx") == [0, 1]
