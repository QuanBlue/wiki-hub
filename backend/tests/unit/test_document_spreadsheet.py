"""Excel/CSV import: workbooks and CSV files into tables."""

from __future__ import annotations

import io
import re
import zipfile
from datetime import date, datetime, time, timedelta
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.core.exceptions import BadRequestError
from app.modules.document_import import spreadsheet
from app.modules.document_import.convert import extract_document, format_for_filename
from app.modules.document_import.spreadsheet import (
    extract_spreadsheet,
    fit_column_widths,
    format_cell_value,
)


def _save(workbook: Workbook, tmp_path: Path, name: str = "Book.xlsx") -> Path:
    path = tmp_path / name
    workbook.save(path)
    return path


def _plain(html: str) -> str:
    """The table without its fitted column widths, which have tests of their own."""
    return re.sub(r' colwidth="[^"]*"', "", html)


def _xlsx_bytes(workbook: Workbook) -> bytes:
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


class TestFormats:
    def test_spreadsheet_extensions_are_importable(self):
        assert format_for_filename("Budget.XLSX") == "xlsx"
        assert format_for_filename("people.csv") == "csv"
        # Macro-enabled workbooks stay out, like `.docm`.
        assert format_for_filename("macro.xlsm") is None


class TestNumberFormats:
    @pytest.mark.parametrize(
        ("value", "number_format", "expected"),
        [
            (None, "General", ""),
            (True, "General", "TRUE"),
            (3.0, "General", "3"),
            (0.1 + 0.2, "General", "0.3"),
            (1234567.5, "#,##0.00", "1,234,567.50"),
            (1234567, "#,##0", "1,234,567"),
            (0.125, "0.0%", "12.5%"),
            (-0.05, "0%", "-5%"),
            (2000000, "#,##0 [$₫-42A]", "2,000,000 ₫"),
            (12.5, '"$"#,##0.00', "$12.50"),
            (3.14159, "0.00", "3.14"),
            (42, "@", "42"),
        ],
    )
    def test_numbers_read_as_the_sheet_shows_them(self, value, number_format, expected):
        assert format_cell_value(value, number_format) == expected

    @pytest.mark.parametrize(
        ("value", "number_format", "expected"),
        [
            (datetime(2026, 1, 31), "dd/mm/yyyy", "31/01/2026"),
            (datetime(2026, 1, 31), "m/d/yyyy", "1/31/2026"),
            (datetime(2026, 2, 28, 14, 30), "dd/mm/yyyy hh:mm", "28/02/2026 14:30"),
            (datetime(2026, 2, 28, 14, 5), "h:mm AM/PM", "2:05 PM"),
            (datetime(2026, 3, 1), "d mmm yyyy", "1 Mar 2026"),
            (datetime(2026, 3, 1), "General", "2026-03-01"),
            (date(2026, 3, 1), "yyyy-mm-dd", "2026-03-01"),
            (time(9, 7, 3), "hh:mm:ss", "09:07:03"),
        ],
    )
    def test_dates_follow_the_cell_format(self, value, number_format, expected):
        assert format_cell_value(value, number_format) == expected


class TestXlsx:
    def test_a_sheet_becomes_a_table_with_a_header_row(self, tmp_path):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Name", "Age"])
        sheet.append(["An", 30])
        html = _plain(extract_spreadsheet(_save(workbook, tmp_path), doc_format="xlsx").html)

        assert html.startswith("<table><thead><tr><th><p>Name</p></th>")
        assert "<td><p>An</p></td><td><p>30</p></td>" in html
        # One sheet needs no heading naming it.
        assert "<h2>" not in html

    def test_formatting_merges_and_links_are_carried_over(self, tmp_path):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Item", "Note"])
        sheet["A2"] = "Total"
        sheet["A2"].font = Font(bold=True, italic=True, color="FFCC0000")
        sheet["A2"].fill = PatternFill("solid", fgColor="FFDDEEFF")
        sheet["A2"].alignment = Alignment(horizontal="center")
        sheet.merge_cells("A2:B2")
        sheet["A3"] = "Docs"
        sheet["A3"].hyperlink = "https://example.com/docs"
        sheet["B3"] = "line one\nline two"
        sheet.column_dimensions["A"].width = 20
        html = _plain(extract_spreadsheet(_save(workbook, tmp_path), doc_format="xlsx").html)

        assert 'colspan="2"' in html
        assert "background-color: #ddeeff" in html
        assert "color: #cc0000" in html
        assert "text-align: center" in html
        assert "<strong><em>Total</em></strong>" in html
        assert '<a href="https://example.com/docs">Docs</a>' in html
        assert "line one<br>line two" in html
        # The cell under the merge is not rendered as a stray extra column.
        assert html.count("<tr>") == 3

    def test_cell_text_is_escaped(self, tmp_path):
        workbook = Workbook()
        workbook.active.append(["<script>alert(1)</script>"])
        html = _plain(extract_spreadsheet(_save(workbook, tmp_path), doc_format="xlsx").html)
        assert "<script>" not in html
        assert "&lt;script&gt;" in html

    def test_unsafe_link_schemes_are_dropped(self, tmp_path):
        workbook = Workbook()
        workbook.active["A1"] = "click"
        workbook.active["A1"].hyperlink = "javascript:alert(1)"
        html = _plain(extract_spreadsheet(_save(workbook, tmp_path), doc_format="xlsx").html)
        assert "javascript" not in html
        assert "<a " not in html

    def test_several_sheets_become_sections_and_hidden_ones_are_reported(self, tmp_path):
        workbook = Workbook()
        workbook.active.title = "Q1"
        workbook.active["A1"] = "one"
        workbook.create_sheet("Empty")
        hidden = workbook.create_sheet("Secret")
        hidden["A1"] = "x"
        hidden.sheet_state = "hidden"
        workbook.create_sheet("Q2")["C3"] = "two"
        result = extract_spreadsheet(_save(workbook, tmp_path), doc_format="xlsx")

        # One page per sheet with something in it, in the workbook's order;
        # the file's own page is built from them by the worker.
        assert result.html == ""
        assert [title for title, _ in result.sections] == ["Q1", "Q2"]
        assert any("Secret" in warning for warning in result.warnings)
        # The table starts at the first used cell, not at A1.
        assert _plain(result.sections[1][1]) == (
            "<table><tbody><tr><td><p>two</p></td></tr></tbody></table>"
        )

    def test_one_sheet_with_data_stays_one_page(self, tmp_path):
        workbook = Workbook()
        workbook.active["A1"] = "only"
        workbook.create_sheet("Blank")
        result = extract_spreadsheet(_save(workbook, tmp_path), doc_format="xlsx")
        assert result.sections == []
        assert "only" in result.html

    def test_sheets_past_the_cap_are_reported(self, tmp_path, monkeypatch):
        monkeypatch.setattr(spreadsheet, "MAX_SHEETS", 2)
        workbook = Workbook()
        workbook.active["A1"] = "a"
        workbook.create_sheet("B")["A1"] = "b"
        workbook.create_sheet("C")["A1"] = "c"
        result = extract_spreadsheet(_save(workbook, tmp_path), doc_format="xlsx")
        assert len(result.sections) == 2
        assert any('"C"' in warning for warning in result.warnings)

    def test_a_large_sheet_is_cut_with_a_warning(self, tmp_path, monkeypatch):
        monkeypatch.setattr(spreadsheet, "MAX_CELLS", 20)
        workbook = Workbook()
        for index in range(30):
            workbook.active.append([index, index * 2])
        result = extract_spreadsheet(_save(workbook, tmp_path), doc_format="xlsx")
        assert result.html.count("<tr>") == 10
        assert any("first 10 rows" in warning for warning in result.warnings)

    def test_an_empty_workbook_says_so(self, tmp_path):
        result = extract_spreadsheet(_save(Workbook(), tmp_path), doc_format="xlsx")
        assert result.html == ""
        assert any("no data" in warning for warning in result.warnings)

    def test_a_file_that_is_not_a_workbook_is_refused(self, tmp_path):
        path = tmp_path / "fake.xlsx"
        path.write_bytes(b"\xd0\xcf\x11\xe0 an encrypted OLE container")
        with pytest.raises(BadRequestError, match="password-protected"):
            extract_spreadsheet(path, doc_format="xlsx")

        other = tmp_path / "other.xlsx"
        with zipfile.ZipFile(other, "w") as archive:
            archive.writestr("word/document.xml", "<w:document/>")
        with pytest.raises(BadRequestError, match="not an Excel workbook"):
            extract_spreadsheet(other, doc_format="xlsx")

    @pytest.mark.asyncio
    async def test_extract_document_routes_spreadsheets_here(self, tmp_path):
        workbook = Workbook()
        workbook.active.append(["a", "b"])
        path = _save(workbook, tmp_path)
        result = await extract_document(path, filename=path.name, workdir=tmp_path)
        assert "<table>" in result.html


class TestCsv:
    def test_delimiter_and_encoding_are_detected(self, tmp_path):
        path = tmp_path / "people.csv"
        path.write_bytes('Tên;Tuổi\nAn;30\nBình;"4;5"\n\n'.encode("utf-8-sig"))
        html = _plain(extract_spreadsheet(path, doc_format="csv").html)
        assert "<th><p>Tên</p></th><th><p>Tuổi</p></th>" in html
        assert "<td><p>4;5</p></td>" in html
        assert html.count("<tr>") == 3

    def test_ragged_rows_are_padded(self, tmp_path):
        path = tmp_path / "ragged.csv"
        path.write_text("a,b,c\n1\n", encoding="utf-8")
        html = _plain(extract_spreadsheet(path, doc_format="csv").html)
        assert html.count("<td>") == 3


class TestEdges:
    """Files at the edges: odd encodings, limits, malformed workbooks and rare formats."""

    def test_a_utf16_csv_with_a_byte_order_mark_is_read(self, tmp_path):
        path = tmp_path / "wide.csv"
        path.write_bytes("Tên,Tuổi\nAn,30\n".encode("utf-16"))
        assert "<th><p>Tên</p></th>" in _plain(extract_spreadsheet(path, doc_format="csv").html)

    @pytest.mark.parametrize("data", [b"caf\xe9,x\n1,2\n", b"caf\xe9,xy\n1,2\n"])
    def test_a_legacy_ansi_csv_falls_back_to_cp1252(self, tmp_path, data):
        # Both lengths: an odd one fails as UTF-16 outright, an even one decodes
        # as UTF-16 but has no byte order mark to trust it by.
        path = tmp_path / "legacy.csv"
        path.write_bytes(data)
        assert "café" in _plain(extract_spreadsheet(path, doc_format="csv").html)

    def test_a_csv_over_the_ceiling_is_refused(self, tmp_path, monkeypatch):
        monkeypatch.setattr(spreadsheet.settings, "document_import_max_media_bytes", 4)
        path = tmp_path / "big.csv"
        path.write_text("a,b\n1,2\n", encoding="utf-8")
        with pytest.raises(BadRequestError, match="more content"):
            extract_spreadsheet(path, doc_format="csv")

    def test_a_long_csv_is_cut_with_a_warning(self, tmp_path, monkeypatch):
        monkeypatch.setattr(spreadsheet, "MAX_CELLS", 6)
        path = tmp_path / "long.csv"
        path.write_text("".join(f"{i},{i}\n" for i in range(10)), encoding="utf-8")
        result = extract_spreadsheet(path, doc_format="csv")
        assert result.html.count("<tr>") == 3
        assert any("first 3 rows" in warning for warning in result.warnings)

    def test_an_empty_csv_says_so(self, tmp_path):
        path = tmp_path / "empty.csv"
        path.write_text(",\n,\n", encoding="utf-8")
        result = extract_spreadsheet(path, doc_format="csv")
        assert result.html == ""
        assert result.warnings == ["This file has no data to import."]

    def test_a_workbook_that_unpacks_too_large_is_refused(self, tmp_path, monkeypatch):
        workbook = Workbook()
        workbook.active["A1"] = "x"
        path = _save(workbook, tmp_path)
        monkeypatch.setattr(spreadsheet.settings, "document_import_max_media_bytes", 10)
        with pytest.raises(BadRequestError, match="more content"):
            extract_spreadsheet(path, doc_format="xlsx")

    def test_a_workbook_openpyxl_cannot_parse_is_refused(self, tmp_path):
        path = tmp_path / "broken.xlsx"
        with zipfile.ZipFile(path, "w") as archive:
            archive.writestr("xl/workbook.xml", "not xml at all")
        with pytest.raises(BadRequestError, match="could not be read"):
            extract_spreadsheet(path, doc_format="xlsx")

    def test_a_very_wide_sheet_is_cut_with_a_warning(self, tmp_path, monkeypatch):
        monkeypatch.setattr(spreadsheet, "MAX_COLUMNS", 2)
        workbook = Workbook()
        workbook.active.append(["a", "b", "c", "d"])
        result = extract_spreadsheet(_save(workbook, tmp_path), doc_format="xlsx")
        assert "<p>c</p>" not in result.html
        assert any("first 2 columns" in warning for warning in result.warnings)

    def test_vertical_merges_and_stray_merges_are_handled(self, tmp_path, monkeypatch):
        monkeypatch.setattr(spreadsheet, "MAX_CELLS", 4)
        workbook = Workbook()
        sheet = workbook.active
        for row in (["a", "b"], ["c", "d"], ["e", "f"], ["g", "h"]):
            sheet.append(row)
        sheet.merge_cells("A1:A2")  # spans rows
        sheet.merge_cells("B1:B2")
        sheet.merge_cells("B2:B3")  # overlaps the one above: its anchor is already covered
        sheet.merge_cells("A4:B4")  # beyond the rows the cap leaves
        html = _plain(extract_spreadsheet(_save(workbook, tmp_path), doc_format="xlsx").html)
        assert 'rowspan="2"' in html
        assert "<p>g</p>" not in html

    @pytest.mark.parametrize(
        ("value", "number_format", "expected"),
        [
            (timedelta(hours=26, minutes=3, seconds=9), "[h]:mm:ss", "26:03:09"),
            (5, "text", "5"),
            (2.5, "text", "2.5"),
            (12345.678, "0.00E+00", "1.23E+04"),
            (1234.5, "#,##0.0%", "123,450.0%"),
            (datetime(2026, 3, 9, 7, 5, 4), "yy-m-d h:m:s a/p", "26-3-9 7:5:4 A"),
            (datetime(2026, 3, 9, 19, 5), "dddd d mmmm", "Monday 9 March"),
            (datetime(2026, 3, 9), 'ddd "the" dd', "Mon the 09"),
            (datetime(2026, 3, 9, 14, 5), "General", "2026-03-09 14:05"),
            (datetime(2026, 3, 9), "@", "2026-03-09"),
            (time(14, 5), "yyyy", "14:05:00"),
            (datetime(2026, 3, 9, 14, 5), 'm "min"', "3 min"),
            # Minutes before seconds, with no hour in front to say so.
            (datetime(2026, 3, 9, 14, 7, 9), "mm:ss", "07:09"),
        ],
    )
    def test_rarer_formats(self, value, number_format, expected):
        assert format_cell_value(value, number_format) == expected


def _widths(html: str) -> list[int]:
    """The fitted width of each column, read off the header row."""
    header = html.split("</tr>", 1)[0]
    return [int(w) for value in re.findall(r'colwidth="([^"]*)"', header) for w in value.split(",")]


class TestColumnWidths:
    """Columns sized by what they hold, not by the sheet's own widths."""

    NOTE = (
        "Cần xác định nhu cầu vận hành ngay từ giai đoạn lập PAKD và đàm phán, ký HĐ; "
        "đưa hạng mục vận hành vào HĐ ngay từ đầu để có cơ sở bố trí nguồn lực."
    )

    def test_a_column_of_paragraphs_gets_the_room_and_a_column_of_codes_stays_narrow(
        self, tmp_path
    ):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["STT", "Đơn vị", "Bài học kinh nghiệm", "Ghi chú"])
        for index in range(1, 4):
            sheet.append([index, "TT ĐTTM", self.NOTE * 2, "x"])
        # Set the way the sheet had them: the notes squeezed, the codes wide.
        sheet.column_dimensions["A"].width = 60
        sheet.column_dimensions["C"].width = 8
        html = extract_spreadsheet(_save(workbook, tmp_path), doc_format="xlsx").html

        number, unit, notes, remark = _widths(html)
        assert notes == max(number, unit, notes, remark)
        assert notes >= 400
        assert number < 100
        assert remark < 100
        # The space is shared out within the page's width.
        assert number + unit + notes + remark <= spreadsheet._TABLE_TARGET_PX + 4

    def test_a_long_header_wraps_instead_of_widening_its_column(self, tmp_path):
        workbook = Workbook()
        workbook.active.append(["Tần suất xuất hiện của lỗi trong các dự án", "Mô tả"])
        workbook.active.append(["Cao", self.NOTE])
        first, _second = _widths(
            extract_spreadsheet(_save(workbook, tmp_path), doc_format="xlsx").html
        )
        # As wide as its longest word or value needs, not as its title.
        assert first < 120

    def test_a_table_that_fits_gets_the_widths_it_would_like(self, tmp_path):
        path = tmp_path / "small.csv"
        path.write_text("Name,Age\nAn,30\n", encoding="utf-8")
        widths = _widths(extract_spreadsheet(path, doc_format="csv").html)
        assert widths == [spreadsheet._chars_px(4), spreadsheet._chars_px(3)]

    def test_too_many_columns_keep_a_readable_line_and_scroll(self, monkeypatch):
        monkeypatch.setattr(spreadsheet, "_TABLE_TARGET_PX", 100)
        grid = [[spreadsheet._Cell(text="longword " * 20) for _ in range(4)] for _ in range(2)]
        readable = spreadsheet._chars_px(spreadsheet._READABLE_CHARS)
        widths = fit_column_widths(grid)
        # Wider than the page altogether: a paragraph is not squeezed to fit,
        # and the table takes a little room beyond the bare minimum.
        assert widths == [round(readable * spreadsheet._SCROLLING_ROOM)] * 4

    def test_merged_and_empty_cells_do_not_size_a_column(self):
        title = spreadsheet._Cell(text="A title spread across both columns of the table")
        title.colspan = 2
        grid = [[title, None], [spreadsheet._Cell(text="a"), spreadsheet._Cell(text=" ")]]
        assert fit_column_widths(grid) == [spreadsheet._chars_px(1), spreadsheet._MIN_COLUMN_PX]

    def test_a_very_long_word_wraps_rather_than_widening_the_minimum(self):
        url = "https://example.com/" + "a" * 200
        grid = [[spreadsheet._Cell(text="Link")], [spreadsheet._Cell(text=url)]]
        assert fit_column_widths(grid) == [spreadsheet._MAX_COLUMN_PX]

    def test_when_the_columns_do_not_all_fit_the_room_goes_to_those_with_most_text(self):
        long, longer = "word " * 40, "word " * 120
        grid = [
            [spreadsheet._Cell(text=t) for t in ("No", "Case", "Lesson", "Fix", "Owner")],
            [spreadsheet._Cell(text=t) for t in ("1", long, longer, longer, "TT ĐTTM")],
        ]
        number, case, lesson, fix, owner = fit_column_widths(grid)
        # Too wide to give everyone what they would like...
        assert 3 * spreadsheet._MAX_COLUMN_PX > spreadsheet._TABLE_TARGET_PX
        # ...so the room is shared: the text columns take it, the most text the most,
        readable = spreadsheet._chars_px(spreadsheet._READABLE_CHARS)
        minimums = spreadsheet._chars_px(2) + 3 * readable + spreadsheet._chars_px(7)
        target = max(spreadsheet._TABLE_TARGET_PX, round(minimums * spreadsheet._SCROLLING_ROOM))
        assert abs((number + case + lesson + fix + owner) - target) <= 4
        assert owner < case < lesson
        assert lesson == fix
        # and a short column keeps at least its longest word.
        assert number >= spreadsheet._chars_px(2)
        assert owner >= spreadsheet._chars_px(4)

    def test_a_few_long_cells_wrap_rather_than_widening_a_column_of_codes(self):
        # Notes above the table, then the table: a column of short codes keeps
        # to the codes' width, and the notes wrap onto two lines.
        notes = [["Gởi kết quả Rà soát/đánh giá", "31/10"], ["Kết quả & kế hoạch khắc phục", "20/11"]]
        codes = [[f"GG-LEG-{n:02}", "Pháp lý"] for n in range(1, 20)]
        grid = [
            [spreadsheet._Cell(text=t) for t in ("Plan", "Date")],
            *[[spreadsheet._Cell(text=t) for t in row] for row in notes + codes],
        ]
        code, _ = fit_column_widths(grid)
        assert code == spreadsheet._chars_px(len("GG-LEG-01"))
        # With only a few cells, each still counts: nothing in them is an outlier.
        few = [[spreadsheet._Cell(text="Code")], *[[spreadsheet._Cell(text=t)] for t in notes[0]]]
        assert fit_column_widths(few) == [spreadsheet._chars_px(len(notes[0][0]))]

    def test_a_column_stops_growing_at_the_width_it_would_like(self):
        # The second column has by far the most text, but in lines of forty
        # characters: it stops at that width and the rest goes to the others.
        line = "word " * 8
        grid = [
            [spreadsheet._Cell(text=t) for t in ("A", "B", "C", "D")],
            [
                spreadsheet._Cell(text=t)
                for t in ("x " * 200, (line.strip() + "\n") * 400, "y " * 150, "z " * 150)
            ],
        ]
        first, second, third, fourth = fit_column_widths(grid)
        assert second == spreadsheet._chars_px(len(line.strip()))
        assert min(first, third, fourth) > spreadsheet._chars_px(spreadsheet._READABLE_CHARS)
