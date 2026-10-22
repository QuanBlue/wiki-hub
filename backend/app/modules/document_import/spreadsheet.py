"""Turn an Excel workbook or a CSV file into tables.

pandoc has no spreadsheet reader, so this module builds the HTML itself. Each
sheet becomes one table. A workbook with one sheet that has something in it is
one page; one with several comes back as ``sections``, one per sheet, which the
worker turns into a page named after the file with a child page per sheet. The
first row is read as the table's header row, which is how nearly every sheet
people import is laid out.

What is carried over is what a wiki table can show: merged cells (as
colspan/rowspan), bold, italic, underline and strike-through, text and fill
colours, horizontal alignment, line breaks inside a cell, hyperlinks, and the
number formats a reader would recognise (dates, percents, thousands
separators, fixed decimals). Column widths are fitted to what each column
holds rather than copied from the sheet (see ``fit_column_widths``), and row
heights follow from them. Formulas are imported as the value the
file last calculated, which is the value its author saw. Charts, images,
comments, conditional formatting and hidden sheets are not imported; a hidden
sheet is named in the import's warnings so its absence is not a surprise.

Every value is escaped here and the result still goes through
``sanitize_imported_html`` afterwards, like every other format.
"""

from __future__ import annotations

import csv
import io
import re
import zipfile
from datetime import date, datetime, time, timedelta
from html import escape
from pathlib import Path
from typing import Any

from app.core.config import settings
from app.core.exceptions import BadRequestError
from app.modules.document_import.convert import DocumentTooComplexError, ExtractedDocument

#: Formats this module reads, keyed the way ``FORMAT_BY_EXTENSION`` names them.
SPREADSHEET_FORMATS = frozenset({"xlsx", "csv"})

#: Cells read per sheet (each sheet is its own page). A page this size is
#: already past what anyone reads as a table, and the page-length limit would
#: cut it anyway; the cap keeps a sheet with a million formatted rows from
#: costing minutes of worker time to build HTML that is then thrown away.
MAX_CELLS = 50_000

#: Sheets imported per workbook, each one a page.
MAX_SHEETS = 50

#: Columns read per sheet. Wider than this is a data dump, not a table.
MAX_COLUMNS = 60

#: Column width bounds, in pixels, for the editor's resizable columns. The
#: minimum is the one the page reader gives every cell (`min-w-24` in
#: `readerClassName`), so the widths fitted here add up to what is shown.
_MIN_COLUMN_PX = 96
_MAX_COLUMN_PX = 480
#: Roughly one character of the page's body text, and a cell's own padding.
_CHAR_PX = 7.5
_CELL_PADDING_PX = 24
#: The width a table is fitted into: a full-width page's content column. A
#: table whose columns cannot fit even at their narrowest scrolls sideways.
_TABLE_TARGET_PX = 1100
#: A longer "word" (a URL, a long code) wraps rather than widening its column.
_MAX_WORD_CHARS = 24
#: A column of sentences is never narrower than about this many characters a
#: line, however many columns there are: below it a paragraph turns into a
#: tower of two-word lines. A table that cannot fit then scrolls sideways,
#: which reads far better.
_READABLE_CHARS = 32
#: How much wider than its columns' minimums a table that has to scroll anyway
#: is allowed to be, so its text columns are not all left at the bare minimum.
_SCROLLING_ROOM = 1.3
#: A column's width follows its usual cell, not its longest: the longest this
#: share of its cells wrap instead. Below ``_FEW_CELLS`` cells, all count.
_OUTLIER_SHARE = 0.15
_FEW_CELLS = 8

_HEX_RGB = re.compile(r"^(?:[0-9A-Fa-f]{2})?([0-9A-Fa-f]{6})$")
_LINK_SCHEMES = ("http://", "https://", "mailto:")


def extract_spreadsheet(source: Path, *, doc_format: str) -> ExtractedDocument:
    """Read ``source`` (``xlsx`` or ``csv``) into HTML tables. Synchronous; run it in a thread."""
    if doc_format == "csv":
        return _extract_csv(source)
    return _extract_xlsx(source)


# --------------------------------------------------------------------------
# CSV
# --------------------------------------------------------------------------


def _decode_csv(data: bytes) -> str:
    # Excel writes UTF-8 with a BOM when asked for "CSV UTF-8", and the local
    # ANSI code page otherwise; cp1252 is the commonest of those and never
    # fails to decode, so it is the last resort rather than an error.
    for encoding in ("utf-8-sig", "utf-16"):
        try:
            text = data.decode(encoding)
        except UnicodeDecodeError:
            continue
        if encoding == "utf-16" and not data.startswith((b"\xff\xfe", b"\xfe\xff")):
            continue
        return text
    return data.decode("cp1252", errors="replace")


def _extract_csv(source: Path) -> ExtractedDocument:
    if source.stat().st_size > settings.document_import_max_media_bytes:
        raise DocumentTooComplexError("This file has more content than WikiHub can process.")
    text = _decode_csv(source.read_bytes())
    sample = text[:16_384]
    try:
        dialect: Any = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows: list[list[str]] = []
    warnings: list[str] = []
    cells = 0
    for row in csv.reader(io.StringIO(text), dialect):
        if cells + len(row) > MAX_CELLS:
            warnings.append(
                f"Only the first {len(rows)} rows were imported; the file is larger "
                "than one page can hold."
            )
            break
        rows.append(row[:MAX_COLUMNS])
        cells += len(row)
    while rows and not any(value.strip() for value in rows[-1]):
        rows.pop()
    if not rows:
        return ExtractedDocument(html="", warnings=["This file has no data to import."])
    width = max(len(row) for row in rows)
    grid: list[list[_Cell | None]] = [
        [*(_Cell(text=value) for value in row), *(_Cell() for _ in range(width - len(row)))]
        for row in rows
    ]
    html = _render_table(grid, widths=fit_column_widths(grid))
    return ExtractedDocument(html=html, warnings=warnings)


# --------------------------------------------------------------------------
# XLSX
# --------------------------------------------------------------------------


class _Cell:
    """One cell as the table will show it."""

    __slots__ = (
        "align",
        "bold",
        "colspan",
        "fill",
        "font_color",
        "italic",
        "link",
        "rowspan",
        "strike",
        "text",
        "underline",
    )

    def __init__(self, text: str = "") -> None:
        self.text = text
        self.bold = self.italic = self.underline = self.strike = False
        self.font_color: str | None = None
        self.fill: str | None = None
        self.align: str | None = None
        self.link: str | None = None
        self.colspan = 1
        self.rowspan = 1


#: A cell covered by a merge; it is not rendered at all.
_COVERED = None


def _check_archive(source: Path) -> None:
    """Refuse a file that is not a workbook, or one that unpacks past the media ceiling."""
    try:
        with zipfile.ZipFile(source) as archive:
            unpacked = sum(info.file_size for info in archive.infolist())
            if "xl/workbook.xml" not in archive.namelist():
                raise BadRequestError("This file is not an Excel workbook.")
    except zipfile.BadZipFile as exc:
        # An encrypted workbook is an OLE container rather than a zip, which
        # is the usual way to land here with a file that opens fine in Excel.
        raise BadRequestError(
            "This workbook could not be read. It may be corrupt or password-protected."
        ) from exc
    if unpacked > settings.document_import_max_media_bytes:
        raise DocumentTooComplexError("This workbook has more content than WikiHub can process.")


def _extract_xlsx(source: Path) -> ExtractedDocument:
    import openpyxl  # Imported here: only an import worker ever needs it.

    _check_archive(source)
    try:
        # `data_only` reads the value each formula last calculated rather than
        # the formula text, which is what a reader of the page wants to see.
        # Not `read_only`: that mode cannot see merged cells.
        workbook = openpyxl.load_workbook(source, data_only=True, keep_links=False)
    except Exception as exc:  # openpyxl raises a zoo of types for a bad file
        raise BadRequestError(
            "This workbook could not be read. It may be corrupt or password-protected."
        ) from exc

    warnings: list[str] = []
    sections: list[tuple[str, str]] = []
    try:
        for sheet in workbook.worksheets:
            if sheet.sheet_state != "visible":
                warnings.append(f"The hidden sheet \"{sheet.title}\" was not imported.")
                continue
            if len(sections) >= MAX_SHEETS:
                warnings.append(f"The sheet \"{sheet.title}\" was not imported: only the "
                                f"first {MAX_SHEETS} sheets of a workbook are.")
                continue
            table, _used = _sheet_table(sheet, budget=MAX_CELLS, warnings=warnings)
            if table:
                sections.append((sheet.title, table))
    finally:
        workbook.close()

    if not sections:
        warnings.append("This workbook has no data to import.")
        return ExtractedDocument(html="", warnings=warnings)
    if len(sections) == 1:
        return ExtractedDocument(html=sections[0][1], warnings=warnings)
    return ExtractedDocument(html="", sections=sections, warnings=warnings)


def _used_bounds(sheet: Any) -> tuple[int, int, int, int] | None:
    """The smallest rectangle holding every non-empty cell, as 1-based (top, left, bottom, right).

    ``max_row``/``max_column`` are not good enough on their own: a column
    formatted all the way down, or a stray border, makes them report a
    thousand rows of nothing.
    """
    top = left = None
    bottom = right = 0
    for row in sheet.iter_rows():
        for cell in row:
            value = cell.value
            if value is None or (isinstance(value, str) and not value.strip()):
                continue
            r, c = cell.row, cell.column
            top = r if top is None else min(top, r)
            left = c if left is None else min(left, c)
            bottom, right = max(bottom, r), max(right, c)
    if top is None or left is None:
        return None
    # A merge whose anchor holds the value still spans the cells to its right
    # and below; those belong inside the table too.
    for merged in sheet.merged_cells.ranges:
        if top <= merged.min_row <= bottom and left <= merged.min_col <= right:
            bottom, right = max(bottom, merged.max_row), max(right, merged.max_col)
    return top, left, bottom, right


def _sheet_table(sheet: Any, *, budget: int, warnings: list[str]) -> tuple[str, int]:
    bounds = _used_bounds(sheet)
    if bounds is None:
        return "", 0
    top, left, bottom, right = bounds
    if right - left + 1 > MAX_COLUMNS:
        warnings.append(f"Only the first {MAX_COLUMNS} columns of \"{sheet.title}\" were imported.")
        right = left + MAX_COLUMNS - 1
    width = right - left + 1
    max_rows = max(1, budget // width)
    if bottom - top + 1 > max_rows:
        warnings.append(
            f"Only the first {max_rows} rows of \"{sheet.title}\" were imported; "
            "the sheet is larger than one page can hold."
        )
        bottom = top + max_rows - 1

    grid: list[list[_Cell | None]] = []
    for row in sheet.iter_rows(min_row=top, max_row=bottom, min_col=left, max_col=right):
        grid.append([_read_cell(cell) for cell in row])

    for merged in sheet.merged_cells.ranges:
        r0, c0 = merged.min_row - top, merged.min_col - left
        if not (0 <= r0 < len(grid) and 0 <= c0 < width):
            continue
        r1 = min(merged.max_row - top, len(grid) - 1)
        c1 = min(merged.max_col - left, width - 1)
        anchor = grid[r0][c0]
        if anchor is None:
            continue
        anchor.rowspan, anchor.colspan = r1 - r0 + 1, c1 - c0 + 1
        for r in range(r0, r1 + 1):
            for c in range(c0, c1 + 1):
                if (r, c) != (r0, c0):
                    grid[r][c] = _COVERED

    return _render_table(grid, widths=fit_column_widths(grid)), len(grid) * width


def _chars_px(chars: int) -> int:
    return max(_MIN_COLUMN_PX, min(_MAX_COLUMN_PX, round(chars * _CHAR_PX + _CELL_PADDING_PX)))


def fit_column_widths(grid: list[list[_Cell | None]]) -> list[int]:
    """Column widths, in pixels, that suit what each column holds.

    The sheet's own widths are not used: they were set for a spreadsheet,
    where a cell's text runs over into the next empty cell and rows are a
    fixed height, so a column of long notes is often left narrow and a column
    of short codes wide. On a page that leaves paragraphs one word wide.

    Each column needs at least its longest word (so no word is broken) - and
    a column of sentences at least a readable line - and would like its
    longest line. When every column fits the page at the width it would like,
    that is what it gets. Otherwise the table is given the page's width, or a
    little more than its columns' minimums when those alone are wider than
    the page (it then scrolls sideways), and the room beyond the minimums is
    shared out by how much text each column holds: a column of paragraphs
    takes the room, a column of numbers stays narrow, and none grows past the
    width it would like.

    The header row counts only towards the minimum: a long title over a column
    of short values should wrap, not widen the column. Cells merged across
    columns are left out, since their text belongs to no one column.

    A column is as wide as its cells usually are, not as its longest one: in
    a column of codes ("GG-LEG-01"), the few notes above the table ("Gởi kết
    quả Rà soát/đánh giá") or a second header row wrap onto two lines rather
    than widening every row of codes to their length.
    """
    columns = max((len(row) for row in grid), default=0)
    least = [0] * columns
    titles = [0] * columns
    lines: list[list[int]] = [[] for _ in range(columns)]
    volume = [0] * columns
    for index, row in enumerate(grid):
        for column, cell in enumerate(row):
            if cell is None or cell.colspan > 1 or not cell.text.strip():
                continue
            longest_word = max(len(word) for word in cell.text.split())
            least[column] = max(least[column], min(longest_word, _MAX_WORD_CHARS))
            if index == 0:
                titles[column] = max(titles[column], longest_word)
                continue
            lines[column].append(max(len(line) for line in cell.text.split("\n")))
            volume[column] += len(cell.text)
    most = [max(title, _usual(found)) for title, found in zip(titles, lines, strict=True)]
    minimum = [
        _chars_px(max(word, min(line, _READABLE_CHARS)))
        for word, line in zip(least, most, strict=True)
    ]
    wanted = [max(low, _chars_px(chars)) for low, chars in zip(minimum, most, strict=True)]
    target = max(_TABLE_TARGET_PX, round(sum(minimum) * _SCROLLING_ROOM))
    if sum(wanted) <= target:
        return wanted

    # Shared out by volume, a pass at a time: a column that reaches the width
    # it would like stops there, and what it did not need goes round again.
    widths = [float(low) for low in minimum]
    growing = {c for c in range(columns) if volume[c] and wanted[c] > minimum[c]}
    spare = float(target - sum(minimum))
    while spare >= 1 and growing:
        total = sum(volume[c] for c in growing)
        given = 0.0
        for column in sorted(growing):
            room = wanted[column] - widths[column]
            share = min(room, spare * volume[column] / total)
            widths[column] += share
            given += share
            if share >= room:
                growing.discard(column)
        spare -= given
    return [round(width) for width in widths]


def _usual(lengths: list[int]) -> int:
    """The line length most of a column's cells stay within: all but the longest few.

    With a handful of cells every one counts (a short table has no outliers to
    speak of); with more, the longest ``_OUTLIER_SHARE`` of them are let wrap.
    """
    if not lengths:
        return 0
    ordered = sorted(lengths)
    if len(ordered) < _FEW_CELLS:
        return ordered[-1]
    return ordered[min(len(ordered) - 1, int(len(ordered) * (1 - _OUTLIER_SHARE)))]


def _rgb(color: Any) -> str | None:
    """``#rrggbb`` for an explicit RGB colour; None for theme, indexed or unset colours."""
    if color is None or getattr(color, "type", None) != "rgb":
        return None
    raw = color.rgb if isinstance(color.rgb, str) else None
    match = _HEX_RGB.match(raw or "")
    return f"#{match.group(1).lower()}" if match else None


def _read_cell(cell: Any) -> _Cell:
    result = _Cell(text=format_cell_value(cell.value, cell.number_format or "General"))
    font = cell.font
    if font is not None:
        result.bold = bool(font.b)
        result.italic = bool(font.i)
        result.underline = bool(font.u) and font.u != "none"
        result.strike = bool(font.strike)
        color = _rgb(font.color)
        # Black is every cell's default; writing it out would pin the text
        # black in dark mode too.
        result.font_color = color if color not in {"#000000"} else None
    fill = cell.fill
    if fill is not None and getattr(fill, "fill_type", None) == "solid":
        color = _rgb(fill.fgColor)
        result.fill = color if color not in {"#ffffff"} else None
    alignment = cell.alignment
    horizontal = getattr(alignment, "horizontal", None) if alignment is not None else None
    result.align = {
        "left": "left",
        "center": "center",
        "centerContinuous": "center",
        "right": "right",
        "justify": "justify",
    }.get(horizontal or "")
    hyperlink = cell.hyperlink
    target = getattr(hyperlink, "target", None) if hyperlink is not None else None
    if isinstance(target, str) and target.lower().startswith(_LINK_SCHEMES):
        result.link = target
    return result


# --------------------------------------------------------------------------
# Rendering
# --------------------------------------------------------------------------


def _cell_html(cell: _Cell, tag: str, width: list[int]) -> str:
    content = "<br>".join(escape(line) for line in cell.text.split("\n"))
    if cell.link and content:
        content = f'<a href="{escape(cell.link, quote=True)}">{content}</a>'
    for flag, mark in (
        (cell.strike, "s"),
        (cell.underline, "u"),
        (cell.italic, "em"),
        (cell.bold and tag == "td", "strong"),
    ):
        if flag and content:
            content = f"<{mark}>{content}</{mark}>"
    styles = []
    if cell.fill:
        styles.append(f"background-color: {cell.fill}")
    if cell.font_color:
        styles.append(f"color: {cell.font_color}")
    if cell.align:
        styles.append(f"text-align: {cell.align}")
    attributes = ""
    if cell.colspan > 1:
        attributes += f' colspan="{cell.colspan}"'
    if cell.rowspan > 1:
        attributes += f' rowspan="{cell.rowspan}"'
    if width:
        attributes += f' colwidth="{",".join(str(w) for w in width)}"'
    if styles:
        attributes += f' style="{"; ".join(styles)}"'
    return f"<{tag}{attributes}><p>{content}</p></{tag}>"


def _render_table(grid: list[list[_Cell | None]], *, widths: list[int]) -> str:
    rows: list[str] = []
    for index, row in enumerate(grid):
        tag = "th" if index == 0 and len(grid) > 1 else "td"
        cells = []
        for column, cell in enumerate(row):
            if cell is None:
                continue
            cells.append(_cell_html(cell, tag, widths[column : column + cell.colspan]))
        rows.append(f"<tr>{''.join(cells)}</tr>")
    if len(rows) > 1:
        return f"<table><thead>{rows[0]}</thead><tbody>{''.join(rows[1:])}</tbody></table>"
    return f"<table><tbody>{''.join(rows)}</tbody></table>"


# --------------------------------------------------------------------------
# Number formats
# --------------------------------------------------------------------------

#: `[Red]`, `[$-409]`, `[$€-x-euro]` and other bracketed directives, plus the
#: `_x` / `*x` padding codes; none of them changes the digits shown.
_FORMAT_NOISE = re.compile(r"\[[^\]]*\]|_.|\*.|\\")
_CURRENCY = re.compile(r"\[\$([^\-\]]*)[^\]]*\]")
_NUMBER_PARTS = re.compile(r"^(?P<prefix>[^0#?.,]*)(?P<body>[0#?.,%Ee+\-]*)(?P<suffix>.*)$")
_DATE_TOKEN = re.compile(
    r'"[^"]*"|yyyy|yy|mmmm|mmm|mm|m|dddd|ddd|dd|d|hh|h|ss|s|am/pm|a/p|.', re.IGNORECASE
)
_MONTHS = ("January", "February", "March", "April", "May", "June", "July", "August",
           "September", "October", "November", "December")
_DAYS = ("Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday")


def format_cell_value(value: Any, number_format: str) -> str:
    """The text Excel would show for ``value`` under ``number_format``, as near as is useful."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (datetime, date, time)):
        return _format_date(value, number_format)
    if isinstance(value, timedelta):
        total = int(value.total_seconds())
        return f"{total // 3600}:{total % 3600 // 60:02d}:{total % 60:02d}"
    if isinstance(value, (int, float)):
        return _format_number(value, number_format)
    return str(value)


def _format_number(value: float, number_format: str) -> str:
    # A negative number has its own section when the format gives one
    # (`#,##0;(#,##0)`); the first section with a minus sign is close enough.
    section = number_format.split(";", 1)[0]
    # `[$₫-42A]` is how a currency symbol is stored; keep the symbol, then
    # drop every other bracketed directive and padding code.
    section = _CURRENCY.sub(r"\1", section)
    section = _FORMAT_NOISE.sub("", section)
    section = re.sub(r'"([^"]*)"', r"\1", section)
    if section.strip().lower() in {"", "general", "@"}:
        if isinstance(value, float) and value.is_integer() and abs(value) < 1e15:
            return str(int(value))
        return f"{value:.10g}" if isinstance(value, float) else str(value)
    parts = _NUMBER_PARTS.match(section)
    if parts is None or not re.search(r"[0#?]", parts.group("body")):
        return f"{value:.10g}" if isinstance(value, float) else str(value)
    body = parts.group("body")
    decimals = 0
    point = body.find(".")
    if point != -1:
        decimals = len(re.match(r"[0#?]*", body[point + 1 :]).group(0))  # type: ignore[union-attr]
    grouped = "," in body.split(".", 1)[0]
    if "%" in body:
        value *= 100
    if "e+" in body.lower():
        text = f"{value:.{decimals}E}"
    else:
        text = f"{value:,.{decimals}f}" if grouped else f"{value:.{decimals}f}"
    if "%" in body:
        text += "%"
    return f"{parts.group('prefix')}{text}{parts.group('suffix')}".strip()


def _format_date(value: datetime | date | time, number_format: str) -> str:
    pattern = _FORMAT_NOISE.sub("", number_format.split(";", 1)[0])
    if not re.search(r"[dmyhs]", pattern, re.IGNORECASE):
        return _iso(value)
    try:
        return _render_date(value, pattern)
    except (AttributeError, ValueError):
        return _iso(value)


def _iso(value: datetime | date | time) -> str:
    if isinstance(value, datetime):
        if value.time() == time(0):
            return value.date().isoformat()
        return value.strftime("%Y-%m-%d %H:%M")
    return value.isoformat()


def _render_date(value: datetime | date | time, pattern: str) -> str:
    tokens = _DATE_TOKEN.findall(pattern)
    twelve_hour = any(token.lower() in {"am/pm", "a/p"} for token in tokens)
    out: list[str] = []
    for index, token in enumerate(tokens):
        lower = token.lower()
        if token.startswith('"'):
            out.append(token[1:-1])
        elif lower == "yyyy":
            out.append(f"{value.year:04d}")  # type: ignore[union-attr]
        elif lower == "yy":
            out.append(f"{value.year % 100:02d}")  # type: ignore[union-attr]
        elif lower in {"m", "mm"} and _is_minute(tokens, index):
            out.append(f"{value.minute:02d}" if lower == "mm" else str(value.minute))  # type: ignore[union-attr]
        elif lower == "mmmm":
            out.append(_MONTHS[value.month - 1])  # type: ignore[union-attr]
        elif lower == "mmm":
            out.append(_MONTHS[value.month - 1][:3])  # type: ignore[union-attr]
        elif lower == "mm":
            out.append(f"{value.month:02d}")  # type: ignore[union-attr]
        elif lower == "m":
            out.append(str(value.month))  # type: ignore[union-attr]
        elif lower == "dddd":
            out.append(_DAYS[value.weekday()])  # type: ignore[union-attr]
        elif lower == "ddd":
            out.append(_DAYS[value.weekday()][:3])  # type: ignore[union-attr]
        elif lower == "dd":
            out.append(f"{value.day:02d}")  # type: ignore[union-attr]
        elif lower == "d":
            out.append(str(value.day))  # type: ignore[union-attr]
        elif lower in {"hh", "h"}:
            hour = value.hour  # type: ignore[union-attr]
            if twelve_hour:
                hour = hour % 12 or 12
            out.append(f"{hour:02d}" if lower == "hh" else str(hour))
        elif lower in {"ss", "s"}:
            out.append(f"{value.second:02d}" if lower == "ss" else str(value.second))  # type: ignore[union-attr]
        elif lower == "am/pm":
            out.append("AM" if value.hour < 12 else "PM")  # type: ignore[union-attr]
        elif lower == "a/p":
            out.append("A" if value.hour < 12 else "P")  # type: ignore[union-attr]
        else:
            out.append(token)
    return "".join(out).strip()


def _is_minute(tokens: list[str], index: int) -> bool:
    """``m`` means minutes right after an hour or right before a second, and months otherwise."""
    for previous in reversed(tokens[:index]):
        lower = previous.lower()
        if lower in {"h", "hh"}:
            return True
        if lower[0].isalpha() and not previous.startswith('"'):
            break
    for following in tokens[index + 1 :]:
        lower = following.lower()
        if lower in {"s", "ss"}:
            return True
        if lower[0].isalpha() and not following.startswith('"'):
            break
    return False
